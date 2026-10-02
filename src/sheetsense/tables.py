"""Load spreadsheets into an in-memory SQLite database, one table per sheet.

Real-world sheets are messy: a title above the headings, money typed as "$1,216.50",
a "Total" row at the bottom, the same name typed in different cases. Each of those
is handled here so the SQL that runs later sees clean, typed columns.
"""

from __future__ import annotations

import csv
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, List, Optional, Sequence

MONEY = re.compile(r"^\(?-?\$?\s*-?[\d,]+(\.\d+)?\)?$")
TOTAL_WORDS = {"total", "totals", "grand total", "sum", "subtotal"}


@dataclass
class Column:
    name: str
    heading: str
    kind: str  # "number", "date" or "text"
    values: List[str] = field(default_factory=list)  # a short list of distinct values, for text columns


@dataclass
class Table:
    name: str
    source: str
    sheet: str
    columns: List[Column]
    rows: int
    skipped_totals: int = 0


def snake(text: str) -> str:
    out = re.sub(r"\W+", "_", str(text).strip().lower()).strip("_")
    return out or "column"


def as_number(value) -> Optional[float]:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str) and MONEY.match(value.strip()):
        text = value.strip()
        negative = text.startswith("(") and text.endswith(")") or "-" in text
        digits = re.sub(r"[^\d.]", "", text)
        if digits in ("", "."):
            return None
        number = float(digits)
        return -number if negative else number
    return None


def as_date(value) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        for pattern in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d %b %Y", "%b %d, %Y"):
            try:
                return datetime.strptime(value.strip(), pattern).date().isoformat()
            except ValueError:
                continue
    return None


def read_rows(path: Path) -> Iterable[tuple[str, List[list]]]:
    """Yield (sheet name, rows) for each sheet of an .xlsx file, or the one table of a .csv."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            yield path.stem, [row for row in csv.reader(f)]
        return
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book.worksheets:
            yield sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()


def find_heading_row(rows: Sequence[list]) -> Optional[int]:
    """The first row, among the first ten, where most cells are text: titles above it are skipped."""
    for i, row in enumerate(rows[:10]):
        cells = [c for c in row if c not in (None, "")]
        texts = [c for c in cells if isinstance(c, str) and as_number(c) is None]
        if len(cells) >= 2 and len(texts) >= max(2, int(0.8 * len(cells))):
            return i
    return None


def is_total_row(row: list) -> bool:
    first = next((c for c in row if c not in (None, "")), None)
    return isinstance(first, str) and first.strip().lower().rstrip(":") in TOTAL_WORDS


def is_code(value) -> bool:
    """Digits with a leading zero ("02134", "007") are a code, not a number: the zero matters."""
    return isinstance(value, str) and len(value.strip()) > 1 and value.strip().isdigit() and value.strip().startswith("0")


def kind_of(values: List) -> str:
    present = [v for v in values if v not in (None, "")]
    if not present:
        return "text"
    if any(is_code(v) for v in present):
        return "text"
    if all(as_number(v) is not None for v in present):
        return "number"
    if all(as_date(v) is not None for v in present):
        return "date"
    return "text"


def unique_names(headings: List[str]) -> List[str]:
    seen: dict[str, int] = {}
    names = []
    for heading in headings:
        name = snake(heading)
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        names.append(name)
    return names


def load(db: sqlite3.Connection, paths: Sequence[Path]) -> List[Table]:
    """Load every sheet of every file. Sheets without a heading row are skipped."""
    tables: List[Table] = []
    taken: set[str] = set()
    for path in paths:
        sheets = list(read_rows(path))
        for sheet, rows in sheets:
            base = snake(path.stem if len(sheets) == 1 else f"{path.stem}_{sheet}")
            table = load_sheet(db, path, sheet, rows, base, taken)
            if table:
                tables.append(table)
                taken.add(table.name)
    return tables


def load_sheet(db: sqlite3.Connection, path: Path, sheet: str, rows: List[list], base: str, taken: set) -> Optional[Table]:
    head = find_heading_row(rows)
    if head is None:
        return None
    width = max(i for i, c in enumerate(rows[head]) if c not in (None, "")) + 1
    headings = [str(c).strip() if c not in (None, "") else f"column {i + 1}" for i, c in enumerate(rows[head][:width])]
    body = [(list(r) + [None] * width)[:width] for r in rows[head + 1:]]
    body = [r for r in body if any(c not in (None, "") for c in r)]
    data = [r for r in body if not is_total_row(r)]
    skipped = len(body) - len(data)

    names = unique_names(headings)
    kinds = [kind_of([r[i] for r in data]) for i in range(width)]
    name, n = base, 2
    while name in taken:
        name, n = f"{base}_{n}", n + 1

    # Text compares ignoring case: 'Unpaid' finds 'unpaid', the way a person reads a sheet.
    declared = ", ".join(f'"{c}" {"REAL" if k == "number" else "TEXT COLLATE NOCASE"}' for c, k in zip(names, kinds))
    db.execute(f'CREATE TABLE "{name}" ({declared})')
    db.executemany(
        f'INSERT INTO "{name}" VALUES ({", ".join("?" * width)})',
        [[clean(v, k) for v, k in zip(r, kinds)] for r in data],
    )
    columns = [Column(c, h, k) for c, h, k in zip(names, headings, kinds)]
    for col in columns:
        if col.kind == "text":
            col.values = short_list(db, name, col.name, len(data))
    return Table(name, path.name, sheet, columns, len(data), skipped)


def clean(value, kind: str):
    if value in (None, ""):
        return None
    if kind == "number":
        return as_number(value)
    if kind == "date":
        return as_date(value)
    return str(value).strip()


def short_list(db: sqlite3.Connection, table: str, column: str, rows: int) -> List[str]:
    """A text column's values when there are only a few (a status, a region): the model then spells them right."""
    found = [r[0] for r in db.execute(f'SELECT DISTINCT "{column}" FROM "{table}" WHERE "{column}" IS NOT NULL LIMIT 9')]
    if 0 < len(found) <= 8 and len(found) * 2 <= rows and all(len(v) <= 30 for v in found):
        return found
    return []


def describe(tables: Sequence[Table]) -> str:
    """The schema as the model sees it."""
    lines = []
    for t in tables:
        extra = f"; {t.skipped_totals} total row(s) left out" if t.skipped_totals else ""
        lines.append(f'Table {t.name}: {t.rows} rows (from {t.source}, sheet "{t.sheet}"){extra}')
        for c in t.columns:
            kind = "date (text yyyy-mm-dd)" if c.kind == "date" else c.kind
            values = f"  values: {', '.join(repr(v) for v in c.values)}" if c.values else ""
            lines.append(f'  {c.name} {kind}  [heading "{c.heading}"]{values}')
    return "\n".join(lines)


def open_db() -> sqlite3.Connection:
    # Used by one caller at a time, but not always from the thread that opened it (the web demo).
    return sqlite3.connect(":memory:", check_same_thread=False)
