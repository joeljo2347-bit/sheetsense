"""Things real files and real setups throw at it."""

import pytest

from sheetsense import guard
from sheetsense.cli import main, table
from sheetsense.model import Ollama, ollama_url
from sheetsense.tables import load, open_db


def test_codes_keep_their_leading_zeros(tmp_path):
    path = tmp_path / "customers.csv"
    path.write_text("Customer,Zip\nA,02134\nB,10001\nC,00501\n")
    db = open_db()
    [t] = load(db, [path])
    assert t.columns[1].kind == "text"
    assert [r[0] for r in db.execute("SELECT zip FROM customers")] == ["02134", "10001", "00501"]


def test_accented_headings_keep_their_letters(tmp_path):
    path = tmp_path / "écoles.csv"
    path.write_bytes("﻿Namé,Élèves\nCafé,12\nÉcole,30\n".encode())
    db = open_db()
    [t] = load(db, [path])
    assert [c.name for c in t.columns] == ["namé", "élèves"]
    assert db.execute('SELECT SUM("élèves") FROM "écoles"').fetchone()[0] == 42


def test_each_sheet_of_a_workbook_is_its_own_table(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "North"
    book.active.append(["Rep", "Sales"])
    book.active.append(["Ana", 10])
    south = book.create_sheet("South")
    south.append(["Rep", "Sales"])
    south.append(["Ben", 20])
    book.create_sheet("Notes").append(["just a note"])  # no headings: skipped
    path = tmp_path / "Regions 2025.xlsx"
    book.save(path)
    assert [t.name for t in load(open_db(), [path])] == ["regions_2025_north", "regions_2025_south"]


@pytest.mark.parametrize("sql", [
    "SELECT * FROM t WHERE note = 'a;b'",
    "SELECT * FROM t WHERE note = 'please delete from the list'",
    "SELECT * FROM t WHERE name = 'O''Brien; Sons'",
])
def test_quoted_text_is_data_not_sql(sql):
    assert guard.refuse(sql) is None


def test_a_semicolon_outside_quotes_still_counts():
    assert guard.refuse("SELECT 'a;b'; DELETE FROM t") is not None


def test_an_empty_result_says_so():
    assert table(guard.Result(["a"], [])).splitlines()[-1] == "(no rows)"


def test_a_cut_result_says_so():
    out = table(guard.Result(["n"], [(i,) for i in range(500)], cut=True), limit=50)
    assert out.splitlines()[-1] == "... more than 500 rows; showing the first 50"
    out = table(guard.Result(["n"], [(i,) for i in range(60)]), limit=50)
    assert out.splitlines()[-1] == "... and 10 more rows"


def test_a_file_without_headings_is_reported(tmp_path, capsys):
    path = tmp_path / "numbers.csv"
    path.write_text("1,2\n3,4\n")
    assert main(["describe", str(path)]) == 1
    assert "No sheet has a row of headings" in capsys.readouterr().out


@pytest.mark.parametrize("host,url", [
    ("", "http://127.0.0.1:11434"),
    ("0.0.0.0:11434", "http://127.0.0.1:11434"),
    ("localhost", "http://localhost:11434"),
    ("http://box:9000/", "http://box:9000"),
    ("https://ollama.example.com", "https://ollama.example.com"),
])
def test_ollama_host_is_read_the_way_ollama_reads_it(host, url):
    assert ollama_url(host) == url


def test_no_ollama_is_a_clear_message_not_a_traceback():
    with pytest.raises(ConnectionError, match="Couldn't reach Ollama"):
        Ollama(url="http://127.0.0.1:9", timeout=2).complete([{"role": "user", "content": "hi"}])
